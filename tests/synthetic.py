"""Synthetic source documents with the real shape of 01 §3 (fixture policy: 01 §10, ADR-020).

Values are invented; element names, attribute names, column headers and row layout follow the
live reads recorded in `docs/01-data-scope.md` §3 (S04, S13, S14). Used by the parse and
mapping tests, by the runtime tests and by `scripts/make_example_fixtures.py`.
"""

from __future__ import annotations

import io
from collections.abc import Iterable, Sequence
from datetime import date, datetime, timedelta
from decimal import Decimal
from zoneinfo import ZoneInfo

import openpyxl

from energy_platform.contracts.intervals import local_day_intervals

PRAGUE = ZoneInfo("Europe/Prague")

# ------------------------------------------------------------------ T1: OTE GetImPricePeriodE

OTE_ITEM = """      <Item>
        <Date>{date}</Date>
        <PeriodResolution>{resolution}</PeriodResolution>
        <PeriodIndex>{index}</PeriodIndex>{emerg}{price}{volume}
      </Item>
"""


def ote_im_price_period_response(
    day: date,
    values: Sequence[tuple[int, str | None, str | None]] | None = None,
    *,
    resolution: str = "PT15M",
    emerg: bool = False,
) -> bytes:
    """``GetImPricePeriodE`` response. ``values`` = (PeriodIndex, Price, Volume); None → absent."""
    if values is None:
        n = len(local_day_intervals(day, resolution))
        values = [(i, f"{170 + i * 0.13:.2f}", f"{120 + i * 0.275:.3f}") for i in range(1, n + 1)]
    items = "".join(
        OTE_ITEM.format(
            date=day.isoformat(),
            resolution=resolution,
            index=index,
            emerg="\n        <Emerg>1</Emerg>" if emerg else "",
            price=f"\n        <Price>{price}</Price>" if price is not None else "",
            volume=f"\n        <Volume>{volume}</Volume>" if volume is not None else "",
        )
        for index, price, volume in values
    )
    return (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<soapenv:Envelope xmlns:soapenv="http://schemas.xmlsoap.org/soap/envelope/">\n'
        "  <soapenv:Body>\n"
        '    <GetImPricePeriodEResponse xmlns="http://www.ote-cr.cz/schema/service/public">\n'
        "    <Result>\n" + items + "    </Result>\n"
        "    </GetImPricePeriodEResponse>\n"
        "  </soapenv:Body>\n"
        "</soapenv:Envelope>\n"
    ).encode("utf-8")


def ote_empty_result() -> bytes:
    return ote_im_price_period_response(date(2026, 9, 17), [])


def soap_fault(code: str = "soap:Server", text: str = "Internal error") -> bytes:
    return (
        '<soap:Envelope xmlns:soap="http://schemas.xmlsoap.org/soap/envelope/"><soap:Body>'
        f"<soap:Fault><faultcode>{code}</faultcode><faultstring>{text}</faultstring></soap:Fault>"
        "</soap:Body></soap:Envelope>"
    ).encode()


# ------------------------------------------------------------------ T3: ČEPS Load

CEPS_ITEM = '        <item date="{date}" value1="{v1}" value2="{v2}" />\n'


def ceps_load_response(
    day: date,
    values: Iterable[tuple[datetime, str, str]] | None = None,
    *,
    hours: int | None = None,
) -> bytes:
    """``Load`` response (QH, AVG, RT). ``values`` = (offset-aware start, value1, value2)."""
    if values is None:
        start = datetime.combine(day, datetime.min.time(), tzinfo=PRAGUE)
        end = start + timedelta(hours=hours) if hours else start + timedelta(days=1)
        vals: list[tuple[datetime, str, str]] = []
        ts = start
        i = 0
        while ts < end:
            vals.append((ts, f"{6382.4 + i * 1.5:.3f}", f"{6279.867 + i * 1.25:.3f}"))
            ts += timedelta(minutes=15)
            i += 1
        values = vals
    items = "".join(CEPS_ITEM.format(date=ts.isoformat(), v1=v1, v2=v2) for ts, v1, v2 in values)
    return (
        '<?xml version="1.0" encoding="utf-8"?>\n'
        '<soap:Envelope xmlns:soap="http://schemas.xmlsoap.org/soap/envelope/">\n'
        "  <soap:Body>\n"
        '    <LoadResponse xmlns="https://www.ceps.cz/CepsData/">\n'
        "      <LoadResult>\n"
        "        <root>\n"
        "        <information>\n"
        f"          <dateFrom>{day.isoformat()}T00:00:00</dateFrom>\n"
        f"          <dateTo>{day.isoformat()}T23:59:59</dateTo>\n"
        "          <agregation>QH</agregation>\n"
        "          <function>AVG</function>\n"
        "          <version>RT</version>\n"
        "        </information>\n"
        "        <series>\n"
        '          <serie id="value1" name="Load including pumping [MW]" />\n'
        '          <serie id="value2" name="Load [MW]" />\n'
        "        </series>\n"
        "        <data>\n" + items + "        </data>\n"
        "        </root>\n"
        "      </LoadResult>\n"
        "    </LoadResponse>\n"
        "  </soap:Body>\n"
        "</soap:Envelope>\n"
    ).encode("utf-8")


# ------------------------------------------------------------------ T2: OTE daily XLSX

XLSX_HEADER = (
    "Period",
    "Time interval",
    "Traded volume (MWh)",
    "Traded volume - purchase (MWh)",
    "Traded volume - sold (MWh)",
    "Average price (EUR/MWh)",
    "Minimal price (EUR/MWh)",
    "Maximal price (EUR/MWh)",
    "Last price (EUR/MWh)",
)


def _label(index: int, day: date) -> str:
    """Display label the way the source renders it: local wall clock, DST repeats included."""
    iv = local_day_intervals(day, "PT15M")[index - 1]
    a, b = iv.start.astimezone(PRAGUE), iv.end.astimezone(PRAGUE)
    return f"{a:%H:%M}-{b:%H:%M}"


def ote_xlsx(
    day: date,
    rows: Sequence[tuple[int, Sequence[Decimal | float | int | str | None]]] | None = None,
    *,
    header: Sequence[str] = XLSX_HEADER,
    filled: int | None = None,
) -> bytes:
    """Daily IM_15MIN file: title row 3, header row 6, periods from row 7, footer after.

    ``rows`` = (Period, the seven numeric cells). Default: every period of the day; ``filled``
    limits how many periods carry numbers (a rolling partial day), the rest are blank.
    """
    n = len(local_day_intervals(day, "PT15M"))
    if rows is None:
        rows = []
        for i in range(1, n + 1):
            if filled is not None and i > filled:
                rows.append((i, [None] * 7))
                continue
            vol = round(120 + i * 0.275, 3)
            rows.append(
                (
                    i,
                    [
                        vol,
                        round(vol * 0.5, 3),
                        round(vol * 0.5, 3),
                        round(170 + i * 0.13, 2),
                        round(165 + i * 0.13, 2),
                        round(175 + i * 0.13, 2),
                        round(171 + i * 0.13, 2),
                    ],
                )
            )
    wb = openpyxl.Workbook()
    ws = wb.active
    if ws is None:
        raise RuntimeError("openpyxl gave no active sheet")
    ws.title = "IM"
    ws.cell(row=3, column=1, value=f"Intraday market results — {day:%d.%m.%Y}")
    for col, text in enumerate(header, start=1):
        ws.cell(row=6, column=col, value=text)
    r = 7
    for period, cells in rows:
        ws.cell(row=r, column=1, value=period)
        ws.cell(row=r, column=2, value=_label(period, day))
        for col, value in enumerate(cells, start=3):
            ws.cell(row=r, column=col, value=None if value is None else value)
        r += 1
    ws.cell(row=r, column=1, value="Source: OTE, a.s. Synthetic fixture shape (01 §10).")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
