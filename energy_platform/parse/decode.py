"""Bytes → ``DecodedDocument`` (ADR-027 §2). The only place third-party parsing libraries run.

* ``xlsx`` — openpyxl, read-only, values only; floats become ``Decimal(repr(x))`` (the cell's
  stored double, never a re-parse of text), dates become ISO text;
* ``xml`` / ``soap`` — lxml with entity resolution and network access off; a SOAP ``Fault`` is
  raised as :class:`SoapFault` so the document is quarantined with that reason;
* ``json`` — the standard library;
* ``html-table`` — lxml.html plus a small CSS-selector subset (tag, ``#id``, ``.class``,
  descendant combinator); anything else fails loudly (Phase 4 widens it if a target needs it);
* ``csv`` — the standard library, header row first.
"""

from __future__ import annotations

import csv
import io
import json
import re
from datetime import date, datetime, time
from decimal import Decimal

import openpyxl
from lxml import etree, html

from energy_platform.contracts.parser import (
    Cell,
    DecodedDocument,
    DecodeKind,
    JsonDocument,
    SheetsDocument,
    TabularDocument,
    XmlDocument,
    XmlElement,
)


class DecodeError(ValueError):
    """The payload is not what ``contract.decode`` says it is; the document is quarantined."""


class SoapFault(DecodeError):
    """A SOAP ``Fault`` carried over HTTP 200 (01 §9 fixture list)."""


_XML_PARSER = etree.XMLParser(
    resolve_entities=False, no_network=True, huge_tree=False, remove_comments=True
)
_MAX_ROWS = 1_000_000


def decode(payload: bytes, kind: DecodeKind) -> DecodedDocument:
    if kind == "xlsx":
        return _decode_xlsx(payload)
    if kind in {"xml", "soap"}:
        return _decode_xml(payload, kind)
    if kind == "json":
        return _decode_json(payload)
    if kind == "html-table":
        raise DecodeError("html-table decoding needs a table selector; use decode_html_table")
    if kind == "csv":
        return _decode_csv(payload)
    raise DecodeError(f"unknown decode kind {kind!r}")


# ------------------------------------------------------------------ xlsx


def _cell(value: object) -> Cell:
    if value is None or isinstance(value, str | bool | int | Decimal):
        return value
    if isinstance(value, float):
        return Decimal(repr(value))
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    return str(value)


def _decode_xlsx(payload: bytes) -> SheetsDocument:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types; every one means "not xlsx"
        raise DecodeError(f"not an xlsx workbook: {exc}") from exc
    sheets: dict[str, tuple[tuple[Cell, ...], ...]] = {}
    for ws in wb.worksheets:
        rows: list[tuple[Cell, ...]] = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= _MAX_ROWS:
                raise DecodeError(f"sheet {ws.title!r} exceeds {_MAX_ROWS} rows")
            rows.append(tuple(_cell(v) for v in row))
        sheets[ws.title] = tuple(rows)
    wb.close()
    return SheetsDocument(sheets=sheets)


# ------------------------------------------------------------------ xml / soap


def _element(node: etree._Element) -> XmlElement:
    children = tuple(_element(c) for c in node if isinstance(c.tag, str))
    text = node.text.strip() if node.text and node.text.strip() else None
    return XmlElement(
        tag=str(node.tag),
        attrib={str(k): str(v) for k, v in node.attrib.items()},
        text=text,
        children=children,
    )


def _decode_xml(payload: bytes, kind: DecodeKind) -> XmlDocument:
    try:
        root = etree.fromstring(payload, parser=_XML_PARSER)  # entities and network off
    except etree.XMLSyntaxError as exc:
        raise DecodeError(f"not well-formed XML: {exc}") from exc
    if root is None:
        raise DecodeError("empty XML document")
    doc = XmlDocument(kind="soap" if kind == "soap" else "xml", root=_element(root))
    if kind == "soap":
        if doc.root.local_name != "Envelope":
            raise DecodeError(f"SOAP root must be Envelope, got {doc.root.local_name!r}")
        fault = doc.root.first("Fault")
        if fault is not None:
            code = fault.first("faultcode") or fault.first("Code")
            reason = fault.first("faultstring") or fault.first("Reason")
            raise SoapFault(
                f"SOAP Fault {_text(code)!r}: {_text(reason)!r}",
            )
    return doc


def _text(el: XmlElement | None) -> str:
    if el is None:
        return ""
    if el.text is not None:
        return el.text
    return " ".join(t for t in (_text(c) for c in el.children) if t)


# ------------------------------------------------------------------ json


def _decode_json(payload: bytes) -> JsonDocument:
    try:
        return JsonDocument(data=json.loads(payload.decode("utf-8")))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise DecodeError(f"not JSON: {exc}") from exc


# ------------------------------------------------------------------ csv


def _decode_csv(payload: bytes, delimiter: str = ",") -> TabularDocument:
    try:
        text = payload.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DecodeError(f"csv is not UTF-8: {exc}") from exc
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = [tuple(c.strip() for c in row) for row in reader if row]
    if not rows:
        raise DecodeError("csv has no header row")
    header = tuple(rows[0])
    return TabularDocument(
        kind="csv", header=header, rows=tuple(tuple(_blank(c) for c in r) for r in rows[1:])
    )


def _blank(cell: str) -> Cell:
    return cell if cell != "" else None


# ------------------------------------------------------------------ html-table

_SIMPLE = re.compile(r"^(?P<tag>[a-zA-Z][a-zA-Z0-9]*)?(?P<id>#[\w-]+)?(?P<classes>(?:\.[\w-]+)*)$")


def css_to_xpath(selector: str) -> str:
    """Translate the supported CSS subset to XPath; anything else is a loud error."""
    steps: list[str] = []
    for token in selector.split():
        m = _SIMPLE.match(token)
        if not m or token in {"", ">", "+", "~"}:
            raise DecodeError(f"unsupported table_selector {selector!r} (tag, #id, .class, ' ')")
        tag = m.group("tag") or "*"
        preds: list[str] = []
        if m.group("id"):
            preds.append(f"@id='{m.group('id')[1:]}'")
        for cls in filter(None, m.group("classes").split(".")):
            preds.append(f"contains(concat(' ', normalize-space(@class), ' '), ' {cls} ')")
        steps.append(tag + "".join(f"[{p}]" for p in preds))
    return "//" + "//".join(steps)


def decode_html_table(payload: bytes, table_selector: str) -> TabularDocument:
    try:
        root = html.fromstring(payload)
    except (etree.ParserError, ValueError) as exc:
        raise DecodeError(f"not HTML: {exc}") from exc
    found = root.xpath(css_to_xpath(table_selector))
    tables = [t for t in found if isinstance(t, etree._Element)] if isinstance(found, list) else []
    if not tables:
        raise DecodeError(f"no element matches table_selector {table_selector!r}")
    table = tables[0]
    header: tuple[str, ...] = ()
    rows: list[tuple[Cell, ...]] = []
    for tr in table.iter("tr"):
        cells = [c for c in tr if c.tag in {"th", "td"}]
        if not cells:
            continue
        texts = [" ".join("".join(str(t) for t in c.itertext()).split()) for c in cells]
        if not header and all(c.tag == "th" for c in cells):
            header = tuple(texts)
            continue
        if not header:
            header = tuple(texts)
            continue
        rows.append(tuple(_blank(t) for t in texts))
    if not header:
        raise DecodeError("table has no rows")
    return TabularDocument(kind="html-table", header=header, rows=tuple(rows))
